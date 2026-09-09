import io
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT


def set_cell_border(cell, **kwargs):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('start', 'top', 'end', 'bottom', 'insideH', 'insideV'):
        if edge in kwargs:
            element = OxmlElement(f'w:{edge}')
            for attr, val in kwargs[edge].items():
                element.set(qn(f'w:{attr}'), str(val))
            tcBorders.append(element)
    tcPr.append(tcBorders)


def add_paragraph_run(paragraph, text, bold=False, size=12, font_name='Arial'):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = font_name
    return run


def generate_docx(minit, kehadiran, seksyen, maklumbalas, pencatat, pengerusi, tarikh_formatted):
    doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1.5

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # ── Title ──
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paragraph_run(title, f"MINIT MESYUARAT {minit['tajuk_mesyuarat']} BIL. {minit['bil']}", bold=True, size=14)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paragraph_run(subtitle, minit['hospital'] or 'HOSPITAL SHAH ALAM', bold=True, size=12)

    doc.add_paragraph()

    # ── Meeting Info ──
    masa_mula = minit['masa_mula'] or ''
    masa_tamat = minit['masa_tamat'] or ''
    waktu_mula = 'petang' if masa_mula > '12:00' else 'pagi'
    waktu_tamat = 'petang' if masa_tamat > '12:00' else 'pagi'

    info_table = doc.add_table(rows=3, cols=3)
    info_data = [
        ('Tarikh', tarikh_formatted),
        ('Masa', f"{masa_mula} {waktu_mula} – {masa_tamat} {waktu_tamat}" if masa_mula else ''),
        ('Tempat', minit['tempat'] or ''),
    ]
    for i, (label, value) in enumerate(info_data):
        row = info_table.rows[i]
        run = row.cells[0].paragraphs[0].add_run(label)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        row.cells[1].text = ':'
        row.cells[2].text = value

    for row in info_table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(12)

    doc.add_paragraph()

    # ── Kehadiran ──
    h = doc.add_paragraph()
    add_paragraph_run(h, 'KEHADIRAN', bold=True, size=12)
    h.paragraph_format.space_after = Pt(6)

    categories = [
        ('HADIR', kehadiran.get('hadir', [])),
        ('TIDAK HADIR DENGAN MAAF', kehadiran.get('tidak_hadir', [])),
        ('TURUT HADIR', kehadiran.get('turut_hadir', [])),
    ]

    for cat_name, members in categories:
        if not members and cat_name != 'HADIR':
            continue

        ch = doc.add_paragraph()
        run = add_paragraph_run(ch, cat_name, bold=True, size=12)
        run.underline = True
        ch.paragraph_format.space_before = Pt(6)

        for i, m in enumerate(members):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)

            text = f"{i + 1}.  {m['jawatan']}"
            add_paragraph_run(p, text, size=12)

            if m['is_pengerusi']:
                p2 = doc.add_paragraph()
                p2.paragraph_format.left_indent = Cm(2)
                add_paragraph_run(p2, 'Pengerusi', bold=True, size=12)

            name_p = doc.add_paragraph()
            name_p.paragraph_format.left_indent = Cm(2)
            add_paragraph_run(name_p, m['nama'], size=12)

            if m.get('wakil_nama'):
                w = doc.add_paragraph()
                w.paragraph_format.left_indent = Cm(2)
                add_paragraph_run(w, f"Diwakili oleh : {m['wakil_nama']}", size=12)

    # Page number
    pn = doc.add_paragraph()
    pn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_paragraph_run(pn, 'M/S | 1', size=10)

    # ── Urus Setia ──
    urus_setia = kehadiran.get('urus_setia', [])
    if urus_setia:
        us_h = doc.add_paragraph()
        run = add_paragraph_run(us_h, 'URUS SETIA', bold=True, size=12)
        run.underline = True
        us_h.paragraph_format.space_before = Pt(6)

        for i, m in enumerate(urus_setia):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            add_paragraph_run(p, f"{i + 1}.  {m['jawatan']}", size=12)

            name_p = doc.add_paragraph()
            name_p.paragraph_format.left_indent = Cm(2)
            add_paragraph_run(name_p, m['nama'], size=12)

        if pencatat:
            pc = doc.add_paragraph()
            pc.paragraph_format.left_indent = Cm(2)
            pc.paragraph_format.space_before = Pt(6)
            add_paragraph_run(pc, 'Pencatat Minit', bold=True, size=12)

    doc.add_paragraph()

    # ── Sections ──
    item_number = 1
    for item in seksyen:
        s = item['seksyen']
        sh = doc.add_paragraph()
        sh.paragraph_format.space_before = Pt(12)
        add_paragraph_run(sh, f"{s['nombor_romawi']}.  {s['tajuk']}", bold=True, size=12)

        for p in item.get('perkara', []):
            pp = doc.add_paragraph()
            pp.paragraph_format.left_indent = Cm(1)
            add_paragraph_run(pp, f"{item_number}.  ", bold=False, size=12)
            add_paragraph_run(pp, p['kandungan'] or '', size=12)

            if p.get('tindakan'):
                tp = doc.add_paragraph()
                tp.paragraph_format.left_indent = Cm(2)
                add_paragraph_run(tp, f"[{p['tindakan'].upper()}]", bold=True, size=10)
                if p.get('pegawai'):
                    add_paragraph_run(tp, f"  Pegawai: {p['pegawai']}", size=10)

            item_number += 1

    doc.add_paragraph()

    # ── Signature ──
    doc.add_paragraph()
    sig_table = doc.add_table(rows=4, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    sig_table.cell(0, 0).text = 'Disediakan Oleh:'
    sig_table.cell(0, 1).text = 'Disahkan Oleh:'

    sig_table.cell(1, 0).text = ''
    sig_table.cell(1, 1).text = ''

    sig_table.cell(2, 0).text = f"____________________\n{pencatat['nama'] if pencatat else ''}\n{pencatat['jawatan'] if pencatat else ''}"
    sig_table.cell(2, 1).text = f"____________________\n{pengerusi['nama'] if pengerusi else ''}\n{pengerusi['jawatan'] + ' (Pengerusi)' if pengerusi else ''}"

    sig_table.cell(3, 0).text = 'Tarikh :'
    sig_table.cell(3, 1).text = 'Tarikh :'

    for row in sig_table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(12)

    for row in sig_table.rows:
        for cell in row.cells:
            cell.paragraphs[0].runs[0].bold = True if row == sig_table.rows[0] else False

    pn2 = doc.add_paragraph()
    pn2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_paragraph_run(pn2, 'M/S | 2', size=10)

    # ── Maklum Balas ──
    if maklumbalas:
        doc.add_page_break()

        mb_title = doc.add_paragraph()
        mb_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_paragraph_run(mb_title,
            f"MAKLUM BALAS MINIT MESYUARAT {minit['tajuk_mesyuarat']} {minit['hospital']} BIL {minit['bil']}",
            bold=True, size=12)

        doc.add_paragraph()

        mb_info = doc.add_table(rows=3, cols=3)
        mb_info_data = [
            ('TARIKH', tarikh_formatted),
            ('MASA', f"{minit['masa_mula'] or ''} – {minit['masa_tamat'] or ''}"),
            ('TEMPAT', minit['tempat'] or ''),
        ]
        for i, (label, value) in enumerate(mb_info_data):
            row = mb_info.rows[i]
            run = row.cells[0].paragraphs[0].add_run(label)
            run.bold = True
            run.font.name = 'Arial'
            run.font.size = Pt(11)
            row.cells[1].text = ':'
            row.cells[2].text = value

        doc.add_paragraph()

        mb_table = doc.add_table(rows=len(maklumbalas) + 1, cols=4)
        mb_table.style = 'Table Grid'
        mb_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ['BIL', 'NO RUJUKAN MINIT', 'PERKARA', 'MAKLUM BALAS\nTindakan/Makluman']
        for i, header in enumerate(headers):
            cell = mb_table.rows[0].cells[i]
            cell.text = header
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.bold = True
                    run.font.name = 'Arial'
                    run.font.size = Pt(10)

        for i, mb in enumerate(maklumbalas):
            row = mb_table.rows[i + 1]
            row.cells[0].text = str(mb['bil'] or i + 1)
            row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row.cells[1].text = mb['no_rujukan'] or ''
            row.cells[2].text = mb['perkara'] or ''
            row.cells[3].text = mb['tindakan_makluman_1'] or ''

            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(11)

        pn3 = doc.add_paragraph()
        pn3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_paragraph_run(pn3, 'M/S | 3', size=10)

    # Save to buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
